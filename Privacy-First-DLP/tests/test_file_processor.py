import unittest
from io import BytesIO

from docx import Document
from pypdf import PdfWriter
from pypdf.generic import DecodedStreamObject, DictionaryObject, NameObject

from modules.file_processor import extract_text_from_file


class FileProcessorTests(unittest.TestCase):
    def test_extracts_utf8_text(self):
        self.assertEqual(extract_text_from_file("sample.txt", "hello".encode()), "hello")

    def test_extracts_csv_cells(self):
        result = extract_text_from_file(
            "sample.csv",
            b"name,email\nAlice,alice@example.com\n",
        )
        self.assertIn("alice@example.com", result)

    def test_extracts_docx_paragraphs_and_table_cells(self):
        buffer = BytesIO()
        document = Document()
        document.add_paragraph("DOCX marker")
        table = document.add_table(rows=1, cols=2)
        table.cell(0, 0).text = "Email"
        table.cell(0, 1).text = "alice@example.com"
        document.save(buffer)

        result = extract_text_from_file("sample.docx", buffer.getvalue())
        self.assertIn("DOCX marker", result)
        self.assertIn("alice@example.com", result)

    def test_extracts_pdf_text(self):
        buffer = BytesIO()
        writer = PdfWriter()
        page = writer.add_blank_page(width=200, height=200)
        font = DictionaryObject({
            NameObject("/Type"): NameObject("/Font"),
            NameObject("/Subtype"): NameObject("/Type1"),
            NameObject("/BaseFont"): NameObject("/Helvetica"),
        })
        resources = DictionaryObject({
            NameObject("/Font"): DictionaryObject({NameObject("/F1"): font})
        })
        stream = DecodedStreamObject()
        stream.set_data(b"BT /F1 12 Tf 20 100 Td (PDF marker) Tj ET")
        page[NameObject("/Resources")] = resources
        page[NameObject("/Contents")] = writer._add_object(stream)
        writer.write(buffer)

        self.assertIn("PDF marker", extract_text_from_file("sample.pdf", buffer.getvalue()))

    def test_rejects_unsupported_extension(self):
        with self.assertRaisesRegex(ValueError, "Unsupported file type"):
            extract_text_from_file("sample.xlsx", b"not supported")


if __name__ == "__main__":
    unittest.main()
