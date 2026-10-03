from pathlib import Path
import io

import pandas as pd
from pypdf import PdfReader
from docx import Document


def extract_text_from_file(file_name: str, file_bytes: bytes) -> str:
    """
    Extract readable text from TXT, PDF, DOCX and CSV files.
    """

    extension = Path(file_name).suffix.lower()

    # TXT
    if extension == ".txt":
        return file_bytes.decode("utf-8", errors="ignore")

    # PDF
    if extension == ".pdf":
        reader = PdfReader(io.BytesIO(file_bytes))

        pages = []

        for page in reader.pages:
            text = page.extract_text() or ""
            pages.append(text)

        return "\n".join(pages)

    # DOCX
    if extension == ".docx":
        document = Document(io.BytesIO(file_bytes))

        paragraphs = [
            paragraph.text
            for paragraph in document.paragraphs
            if paragraph.text.strip()
        ]

        for table in document.tables:
            for row in table.rows:
                cells = [cell.text.strip() for cell in row.cells]
                if any(cells):
                    paragraphs.append("\t".join(cells))

        return "\n".join(paragraphs)

    # CSV
    if extension == ".csv":
        dataframe = pd.read_csv(io.BytesIO(file_bytes))

        return dataframe.to_string(index=False)

    raise ValueError(
        "Unsupported file type. Please upload TXT, PDF, DOCX, or CSV."
    )
